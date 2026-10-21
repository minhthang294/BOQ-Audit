"""Safe, bounded Excel reading and AI explanation for the BOQ Excel explainer."""
from __future__ import annotations

import hashlib
import hmac
import json
import math
import re
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.cell import range_boundaries

from app.core.config import get_settings
from app.models.entities import ExcelWorkbook

MAX_CELL_TEXT = 600
MAX_FORMULA_TEXT = 1600
MAX_PREVIEW_ROWS = 200
MAX_PREVIEW_COLS = 100
SHEET_SAMPLE_ROWS = 15
SHEET_SAMPLE_COLS = 20


class ExcelExplainError(ValueError):
    """User-correctable workbook or selection error."""


class ExcelAIBusy(RuntimeError):
    """The shared AI gateway is already serving another turn."""


class ExcelAIUnavailable(ExcelExplainError):
    """The shared AI gateway is unavailable or misconfigured."""


def excel_path(workbook: ExcelWorkbook) -> Path:
    root = get_settings().excel_dir.resolve()
    path = Path(workbook.file_path).resolve()
    if not path.is_relative_to(root) or not path.is_file():
        raise ExcelExplainError("Tệp Excel không còn tồn tại. Vui lòng tải lại tệp.")
    return path


def _normalise(value: Any) -> Any:
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return f"<{len(value)} bytes>"
    if isinstance(value, str):
        return value[:MAX_CELL_TEXT]
    if value is None or isinstance(value, (bool, int, float)):
        return value
    return str(value)[:MAX_CELL_TEXT]


def _formula_functions(formula: str | None) -> list[str]:
    if not formula:
        return []
    return sorted(set(re.findall(r"\b([A-Z][A-Z0-9.]*)\s*\(", formula.upper())))


def _cell_payload(ws_formula, ws_value, row: int, col: int) -> dict[str, Any]:
    ref = f"{get_column_letter(col)}{row}"
    formula_cell = ws_formula.cell(row=row, column=col)
    value_cell = ws_value.cell(row=row, column=col)
    raw = formula_cell.value
    formula = raw[:MAX_FORMULA_TEXT] if isinstance(raw, str) and raw.startswith("=") else None
    value = value_cell.value
    if formula is None and value is None:
        value = raw
    payload: dict[str, Any] = {
        "ref": ref,
        "row": row,
        "col": col,
        "value": _normalise(value),
        "formula": formula,
        "number_format": formula_cell.number_format,
        "kind": "formula" if formula else ("blank" if value is None else "value"),
        "functions": _formula_functions(formula),
    }
    comment = getattr(formula_cell, "comment", None)
    if comment and comment.text:
        payload["comment"] = str(comment.text)[:500]
    return payload


def _zip_safety(path: Path) -> None:
    settings = get_settings()
    try:
        with zipfile.ZipFile(path) as archive:
            uncompressed = sum(item.file_size for item in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ExcelExplainError("Tệp không phải file Excel .xlsx/.xlsm hợp lệ.") from exc
    limit = settings.excel_explain_max_uncompressed_mb * 1024 * 1024
    if uncompressed > limit:
        raise ExcelExplainError(
            f"File Excel giải nén vượt quá {settings.excel_explain_max_uncompressed_mb} MB. "
            "Vui lòng giảm dung lượng hoặc tách bớt sheet."
        )


def _load(path: Path, data_only: bool, read_only: bool = False):
    try:
        return load_workbook(path, data_only=data_only, read_only=read_only, keep_links=False)
    except Exception as exc:
        raise ExcelExplainError("Không đọc được file Excel. Hãy kiểm tra file có bị khóa hoặc hỏng không.") from exc


def inspect_workbook(path: Path) -> dict[str, Any]:
    settings = get_settings()
    _zip_safety(path)
    workbook = _load(path, data_only=False, read_only=True)
    try:
        if not workbook.sheetnames:
            raise ExcelExplainError("File Excel không có sheet nào.")
        if len(workbook.sheetnames) > settings.excel_explain_max_sheets:
            raise ExcelExplainError(
                f"File có {len(workbook.sheetnames)} sheet, vượt quá giới hạn {settings.excel_explain_max_sheets} sheet."
            )
        sheets = [
            {
                "name": ws.title,
                "max_row": max(1, ws.max_row),
                "max_column": max(1, ws.max_column),
                "hidden": ws.sheet_state != "visible",
            }
            for ws in workbook.worksheets
        ]
        active = workbook.active.title if workbook.active else sheets[0]["name"]
        return {"sheets": sheets, "active_sheet": active}
    finally:
        workbook.close()


def _clamp(value: int, low: int, high: int) -> int:
    return max(low, min(high, value))


def preview_sheet(
    path: Path,
    sheet_name: str,
    max_rows: int = 60,
    max_cols: int = 20,
    row_start: int = 1,
    col_start: int = 1,
) -> dict[str, Any]:
    settings = get_settings()
    _zip_safety(path)
    formulas = _load(path, data_only=False, read_only=True)
    values = _load(path, data_only=True, read_only=True)
    try:
        if sheet_name not in formulas.sheetnames:
            raise ExcelExplainError("Sheet không tồn tại trong file.")
        ws_formula = formulas[sheet_name]
        ws_value = values[sheet_name]
        max_row = max(1, ws_formula.max_row or 1)
        max_col = max(1, ws_formula.max_column or 1)
        row_start = max(1, row_start)
        col_start = max(1, col_start)
        max_rows = _clamp(max_rows, 1, MAX_PREVIEW_ROWS)
        max_cols = _clamp(max_cols, 1, MAX_PREVIEW_COLS)
        row_end = min(max_row, row_start + max_rows - 1)
        col_end = min(max_col, col_start + max_cols - 1)
        if row_end < row_start:
            row_start, row_end = 1, min(max_row, row_start + max_rows - 1)
        if col_end < col_start:
            col_start, col_end = 1, min(max_col, col_start + max_cols - 1)

        # Build a coordinate-keyed map for both workbooks (read-only iter_rows
        # re-streams from the start of the file; doing it once per workbook is
        # the cheap part and avoids O(n^2) random cell access).
        formula_map: dict[tuple[int, int], Any] = {}
        value_map: dict[tuple[int, int], Any] = {}
        formula_format: dict[tuple[int, int], str | None] = {}
        f_rows = list(ws_formula.iter_rows(min_row=row_start, max_row=row_end, min_col=col_start, max_col=col_end))
        v_rows = list(ws_value.iter_rows(min_row=row_start, max_row=row_end, min_col=col_start, max_col=col_end))
        for r_idx, (frow, vrow) in enumerate(zip(f_rows, v_rows)):
            for c_idx, (fcell, vcell) in enumerate(zip(frow, vrow)):
                row = r_idx + row_start
                col = c_idx + col_start
                key = (row, col)
                raw = getattr(fcell, "value", None)
                formula_map[key] = raw
                formula_format[key] = getattr(fcell, "number_format", None)
                if vcell is not None:
                    vv = getattr(vcell, "value", None)
                    if vv is not None:
                        value_map[key] = vv


        rows = []
        for row in range(row_start, row_end + 1):
            cells = []
            for col in range(col_start, col_end + 1):
                key = (row, col)
                raw = formula_map.get(key)
                formula = raw[:MAX_FORMULA_TEXT] if isinstance(raw, str) and raw.startswith("=") else None
                value = value_map.get(key, None)
                if formula is None and value is None:
                    value = raw if raw is not None else None
                kind = "formula" if formula else ("blank" if value is None else "value")
                cells.append({
                    "ref": f"{get_column_letter(col)}{row}",
                    "row": row,
                    "col": col,
                    "value": _normalise(value),
                    "formula": formula,
                    "number_format": formula_format.get(key),
                    "kind": kind,
                })
            rows.append(cells)

        merged = []  # read-only mode does not expose mergeCells; kept for schema compat
        return {
            "sheet": sheet_name,
            "max_row": max_row,
            "max_column": max_col,
            "row_start": row_start,
            "col_start": col_start,
            "row_end": row_end,
            "col_end": col_end,
            "columns": [get_column_letter(col) for col in range(col_start, col_end + 1)],
            "rows": rows,
            "merged": merged,
            "truncated": max_row > row_end or max_col > col_end,
            "freeze_panes": None,
            "auto_filter": None,
            "sheet_state": ws_formula.sheet_state,
        }
    finally:
        formulas.close()
        values.close()


def parse_selection(selection: str, sheet_name: str, scope: str) -> dict[str, Any]:
    if scope == "sheet":
        return {"scope": "sheet", "address": f"{sheet_name}!*", "min_col": 1, "min_row": 1, "max_col": None, "max_row": None}
    cleaned = selection.strip().replace("$", "")
    if "!" in cleaned:
        prefix, cleaned = cleaned.rsplit("!", 1)
        prefix = prefix.strip("'").replace("''", "'")
        if prefix and prefix != sheet_name:
            raise ExcelExplainError("Vùng chọn đang trỏ sang sheet khác. Hãy chọn đúng sheet trước.")
    match = re.fullmatch(r"([A-Za-z]{1,3})(\d+)(?::([A-Za-z]{1,3})(\d+))?", cleaned)
    if not match:
        raise ExcelExplainError("Vùng chọn không hợp lệ. Dùng dạng A1, W170 hoặc A1:C10.")
    min_col = column_index_from_string(match.group(1).upper())
    min_row = int(match.group(2))
    max_col = column_index_from_string((match.group(3) or match.group(1)).upper())
    max_row = int((match.group(4) or match.group(2)))
    min_col, max_col = sorted((min_col, max_col))
    min_row, max_row = sorted((min_row, max_row))
    area = (max_col - min_col + 1) * (max_row - min_row + 1)
    limit = get_settings().excel_explain_max_selection_cells
    if area > limit:
        raise ExcelExplainError(
            f"Vùng chọn có {area:,} ô, vượt quá giới hạn {limit:,} ô. "
            "Hãy chọn vùng nhỏ hơn hoặc chọn cả sheet để AI tóm tắt."
        )
    address = f"{sheet_name}!{get_column_letter(min_col)}{min_row}"
    if max_col != min_col or max_row != min_row:
        address += f":{get_column_letter(max_col)}{max_row}"
    return {
        "scope": "cell" if area == 1 else "range",
        "address": address,
        "min_col": min_col,
        "min_row": min_row,
        "max_col": max_col,
        "max_row": max_row,
        "area": area,
    }


def _extract_refs(formula: str | None, default_sheet: str) -> list[dict[str, Any]]:
    if not formula:
        return []
    refs: list[dict[str, Any]] = []
    seen: set[tuple] = set()
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return refs
    for token in tokens:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        raw = token.value.strip()
        if not raw or "[" in raw or "]" in raw or raw.startswith("#"):
            continue
        sheet = default_sheet
        ref = raw
        if "!" in raw:
            sheet_part, ref = raw.rsplit("!", 1)
            sheet = sheet_part.strip("'").replace("''", "'")
        ref = ref.replace("$", "")
        try:
            min_col, min_row, max_col, max_row = range_boundaries(ref)
        except Exception:
            continue
        if min_col is None or min_row is None or max_col is None or max_row is None:
            kind = "full_column" if min_row is None else "full_row"
            key = (sheet, ref, kind)
            if key not in seen:
                seen.add(key)
                refs.append({"sheet": sheet, "ref": ref, "kind": kind, "source": raw})
            continue
        min_col, max_col = sorted((min_col, max_col))
        min_row, max_row = sorted((min_row, max_row))
        key = (sheet, min_col, min_row, max_col, max_row)
        if key in seen:
            continue
        seen.add(key)
        address = f"{get_column_letter(min_col)}{min_row}"
        if max_col != min_col or max_row != min_row:
            address += f":{get_column_letter(max_col)}{max_row}"
        refs.append(
            {
                "sheet": sheet,
                "ref": address,
                "kind": "range" if address != f"{get_column_letter(min_col)}{min_row}" else "cell",
                "min_col": min_col,
                "min_row": min_row,
                "max_col": max_col,
                "max_row": max_row,
                "source": raw,
            }
        )
    return refs


def _sample_coords(
    min_col: int,
    min_row: int,
    max_col: int,
    max_row: int,
    budget: int,
) -> tuple[list[tuple[int, int]], bool]:
    total = (max_col - min_col + 1) * (max_row - min_row + 1)
    if total <= budget:
        return [(row, col) for row in range(min_row, max_row + 1) for col in range(min_col, max_col + 1)], False
    coords: list[tuple[int, int]] = []
    head_rows = max(1, min(max_row - min_row + 1, budget // max(1, max_col - min_col + 1)))
    for row in range(min_row, min(min_row + head_rows, max_row + 1)):
        for col in range(min_col, max_col + 1):
            coords.append((row, col))
            if len(coords) >= budget:
                return coords, True
    for row in range(max(min_row, max_row - 4), max_row + 1):
        for col in range(min_col, max_col + 1):
            coords.append((row, col))
            if len(coords) >= budget:
                return coords, True
    return coords, True


def _sheet_summary(formulas, ws_formula, ws_value) -> dict[str, Any]:
    merged = [str(item) for item in list(ws_formula.merged_cells.ranges)[:100]]
    tables = []
    try:
        tables = list(ws_formula.tables.keys())
    except Exception:
        pass
    validations = 0
    try:
        validations = len(ws_formula.data_validations.dataValidation)
    except Exception:
        pass
    scan_rows = min(max(1, ws_formula.max_row), 3000)
    scan_cols = min(max(1, ws_formula.max_column), 60)
    formula_examples: list[dict[str, Any]] = []
    formula_count = 0
    scanned = 0
    for row in ws_formula.iter_rows(min_row=1, max_row=scan_rows, min_col=1, max_col=scan_cols):
        for cell in row:
            scanned += 1
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formula_count += 1
                if len(formula_examples) < 40:
                    formula_examples.append(
                        {
                            "ref": cell.coordinate,
                            "formula": cell.value[:MAX_FORMULA_TEXT],
                            "functions": _formula_functions(cell.value),
                        }
                    )
            if scanned >= 50000:
                break
        if scanned >= 50000:
            break
    sample_rows: list[list[dict[str, Any]]] = []
    for row in range(1, min(ws_formula.max_row, SHEET_SAMPLE_ROWS) + 1):
        sample_rows.append(
            [_cell_payload(ws_formula, ws_value, row, col) for col in range(1, min(ws_formula.max_column, SHEET_SAMPLE_COLS) + 1)]
        )
    if ws_formula.max_row > SHEET_SAMPLE_ROWS:
        for row in range(max(SHEET_SAMPLE_ROWS + 1, ws_formula.max_row - 4), ws_formula.max_row + 1):
            sample_rows.append(
                [_cell_payload(ws_formula, ws_value, row, col) for col in range(1, min(ws_formula.max_column, SHEET_SAMPLE_COLS) + 1)]
            )
    named_ranges = []
    try:
        for name, definition in list(formulas.defined_names.items())[:30]:
            named_ranges.append({"name": name, "formula": str(definition.attr_text)[:500]})
    except Exception:
        pass
    return {
        "dimension": ws_formula.calculate_dimension(),
        "max_row": max(1, ws_formula.max_row),
        "max_column": max(1, ws_formula.max_column),
        "merged": merged,
        "tables": tables,
        "data_validation_count": validations,
        "freeze_panes": str(ws_formula.freeze_panes) if ws_formula.freeze_panes else None,
        "auto_filter": ws_formula.auto_filter.ref if ws_formula.auto_filter else None,
        "sheet_state": ws_formula.sheet_state,
        "formula_count_scanned": formula_count,
        "formula_scan_scope": f"{scan_rows} dòng x {scan_cols} cột đầu",
        "formula_examples": formula_examples,
        "sample_rows": sample_rows,
        "sample_note": f"Chỉ trích {SHEET_SAMPLE_ROWS} dòng đầu và 5 dòng cuối, tối đa {SHEET_SAMPLE_COLS} cột.",
        "named_ranges": named_ranges,
    }


def _build_references(
    formulas,
    values,
    start_sheet: str,
    selected_cells: list[dict[str, Any]],
    context_cells: list[dict[str, Any]],
    selected_keys: set[tuple[str, str]],
    focus_row: int | None = None,
    focus_col: int | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], bool]:
    settings = get_settings()
    max_refs = settings.excel_explain_max_reference_cells
    reference_cells: list[dict[str, Any]] = []
    reference_ranges: list[dict[str, Any]] = []
    seen_keys = set(selected_keys)
    seen_ranges: set[tuple] = set()
    truncated = False
    queue: list[dict[str, Any]] = []
    for cell in selected_cells + context_cells:
        for ref in _extract_refs(cell.get("formula"), start_sheet):
            queue.append({**ref, "depth": 0})
    while queue and len(reference_cells) < max_refs:
        ref = queue.pop(0)
        sheet_name = ref["sheet"]
        if sheet_name not in formulas.sheetnames:
            reference_ranges.append({"sheet": sheet_name, "ref": ref.get("ref"), "note": "Sheet không tồn tại trong file."})
            continue
        ws_formula = formulas[sheet_name]
        ws_value = values[sheet_name]
        if ref["kind"] == "full_column":
            min_col = column_index_from_string(ref["ref"].split(":")[0].replace("$", ""))
            max_col = column_index_from_string(ref["ref"].split(":")[1].replace("$", ""))
            if sheet_name == start_sheet and focus_row:
                min_row, max_row = max(1, focus_row - 30), min(ws_formula.max_row, focus_row + 40)
            else:
                min_row, max_row = 1, min(ws_formula.max_row, 30)
        elif ref["kind"] == "full_row":
            min_row = int(ref["ref"].split(":")[0])
            max_row = int(ref["ref"].split(":")[1])
            if sheet_name == start_sheet and focus_col:
                min_col, max_col = max(1, focus_col - 10), min(ws_formula.max_column, focus_col + 15)
            else:
                min_col, max_col = 1, min(ws_formula.max_column, 30)
        else:
            min_col, min_row, max_col, max_row = ref["min_col"], ref["min_row"], ref["max_col"], ref["max_row"]
        min_row = max(1, min_row)
        min_col = max(1, min_col)
        max_row = min(max(1, ws_formula.max_row), max_row)
        max_col = min(max(1, ws_formula.max_column), max_col)
        if min_row > max_row or min_col > max_col:
            continue
        key = (sheet_name, min_col, min_row, max_col, max_row)
        if key in seen_ranges:
            continue
        seen_ranges.add(key)
        budget = min(80, max_refs - len(reference_cells))
        coords, sampled = _sample_coords(min_col, min_row, max_col, max_row, max(1, budget))
        truncated = truncated or sampled
        reference_ranges.append(
            {
                "sheet": sheet_name,
                "ref": ref.get("ref"),
                "source": ref.get("source"),
                "sampled": sampled,
                "note": "Chỉ trích mẫu một phần vì vùng tham chiếu quá lớn." if sampled else None,
            }
        )
        for row, col in coords:
            if len(reference_cells) >= max_refs:
                truncated = True
                break
            payload = _cell_payload(ws_formula, ws_value, row, col)
            cell_key = (sheet_name, payload["ref"])
            if cell_key in seen_keys:
                continue
            if ref["kind"] in {"full_column", "full_row"} and payload["kind"] == "blank":
                continue
            seen_keys.add(cell_key)
            payload["sheet"] = sheet_name
            reference_cells.append(payload)
            if ref.get("depth", 0) < 2 and payload.get("formula"):
                for nested in _extract_refs(payload["formula"], sheet_name):
                    queue.append({**nested, "depth": ref.get("depth", 0) + 1})
    return reference_cells, reference_ranges, truncated


def build_context(
    workbook: ExcelWorkbook,
    sheet_name: str,
    scope: str,
    selection: str,
) -> tuple[dict[str, Any], int, int, bool]:
    settings = get_settings()
    path = excel_path(workbook)
    _zip_safety(path)
    formulas = _load(path, data_only=False)
    values = _load(path, data_only=True)
    try:
        if sheet_name not in formulas.sheetnames:
            raise ExcelExplainError("Sheet không tồn tại trong file.")
        ws_formula = formulas[sheet_name]
        ws_value = values[sheet_name]
        summary = _sheet_summary(formulas, ws_formula, ws_value)
        selected_cells: list[dict[str, Any]] = []
        context_cells: list[dict[str, Any]] = []
        focus_row: int | None = None
        focus_col: int | None = None
        if scope == "sheet":
            selected_cells = [cell for row in summary["sample_rows"] for cell in row]
            selection_info = {
                "scope": "sheet",
                "address": f"{sheet_name}!{summary['dimension']}",
                "cell_count": summary["max_row"] * summary["max_column"],
                "note": summary["sample_note"],
            }
        else:
            bounds = parse_selection(selection, sheet_name, scope)
            focus_row = bounds["min_row"]
            focus_col = bounds["min_col"]
            selection_info = {
                "scope": bounds["scope"],
                "address": bounds["address"],
                "cell_count": bounds["area"],
                "note": None,
            }
            for row in range(bounds["min_row"], bounds["max_row"] + 1):
                for col in range(bounds["min_col"], bounds["max_col"] + 1):
                    selected_cells.append(_cell_payload(ws_formula, ws_value, row, col))
            if bounds["min_row"] > 1:
                header_row = bounds["min_row"] - 1
                for col in range(bounds["min_col"], min(bounds["max_col"], bounds["min_col"] + 29) + 1):
                    context_cells.append(_cell_payload(ws_formula, ws_value, header_row, col))
            if bounds["min_col"] > 1:
                label_col = bounds["min_col"] - 1
                for row in range(bounds["min_row"], min(bounds["max_row"], bounds["min_row"] + 29) + 1):
                    context_cells.append(_cell_payload(ws_formula, ws_value, row, label_col))
        selected_keys = {(sheet_name, cell["ref"]) for cell in selected_cells}
        context_keys = {(sheet_name, cell["ref"]) for cell in context_cells}
        context_cells = [cell for cell in context_cells if (sheet_name, cell["ref"]) not in selected_keys]
        reference_cells, reference_ranges, refs_truncated = _build_references(
            formulas, values, sheet_name, selected_cells, context_cells, selected_keys | context_keys, focus_row, focus_col
        )
        for cell in selected_cells:
            cell["sheet"] = sheet_name
        for cell in context_cells:
            cell["sheet"] = sheet_name
        context = {
            "task": "excel_explain",
            "workbook": {
                "filename": workbook.original_filename,
                "sheets": [
                    {"name": ws.title, "max_row": max(1, ws.max_row), "max_column": max(1, ws.max_column), "hidden": ws.sheet_state != "visible"}
                    for ws in formulas.worksheets
                ],
                "active_sheet": sheet_name,
            },
            "selection": selection_info,
            "sheet": summary,
            "selected_cells": selected_cells,
            "context_cells": context_cells,
            "reference_cells": reference_cells,
            "reference_ranges": reference_ranges,
            "evidence_scope": (
                "Chỉ các ô trong vùng chọn, một hàng/cột tiêu đề gần kề và các ô được tham chiếu trực tiếp/gián tiếp "
                f"tối đa 2 tầng, giới hạn {settings.excel_explain_max_reference_cells} ô tham chiếu. "
                "Không đọc toàn bộ workbook; giá trị công thức có thể là giá trị cache do Excel lưu lần cuối."
            ),
        }
        truncated = refs_truncated or len(selected_cells) < (selection_info.get("cell_count") or 0)
        raw = json.dumps(context, ensure_ascii=False)
        if len(raw) > settings.excel_explain_max_context_chars:
            truncated = True
            while len(raw) > settings.excel_explain_max_context_chars and context["reference_cells"]:
                context["reference_cells"].pop()
                raw = json.dumps(context, ensure_ascii=False)
            while len(raw) > settings.excel_explain_max_context_chars and context["context_cells"]:
                context["context_cells"].pop()
                raw = json.dumps(context, ensure_ascii=False)
            while len(raw) > settings.excel_explain_max_context_chars and len(context["selected_cells"]) > 1:
                context["selected_cells"].pop()
                raw = json.dumps(context, ensure_ascii=False)
            context["evidence_scope"] += " Context đã được rút gọn do giới hạn dung lượng."
        return context, len(selected_cells) + len(context_cells), len(reference_cells), truncated
    finally:
        formulas.close()
        values.close()


def _gateway_reply(context: dict[str, Any], message: str, history: list[dict], user_id: int, code: str) -> str:
    settings = get_settings()
    if not settings.chat_gateway_token:
        raise ExcelAIUnavailable("AI chưa được cấu hình. Vui lòng liên hệ quản trị viên.")
    session_key = hmac.new(
        settings.chat_gateway_token.encode(),
        f"excel:{user_id}:{code}".encode(),
        hashlib.sha256,
    ).hexdigest()
    payload = {
        "session_key": session_key,
        "thread_id": None,
        "context": context,
        "history": history[-6:],
        "message": message[:2000],
    }
    request = Request(
        settings.chat_gateway_url.rstrip("/") + "/reply",
        data=json.dumps(payload, ensure_ascii=False).encode(),
        headers={"Content-Type": "application/json", "Authorization": "Bearer " + settings.chat_gateway_token},
        method="POST",
    )
    try:
        with urlopen(request, timeout=settings.chat_timeout_seconds + 10) as response:
            raw = response.read(100001)
    except HTTPError as exc:
        detail = ""
        try:
            detail = json.loads(exc.read()).get("detail", "")
        except Exception:
            detail = ""
        if exc.code == 409:
            raise ExcelAIBusy("AI đang trả lời một câu hỏi khác. Vui lòng thử lại sau.") from None
        if detail:
            raise ExcelAIUnavailable(f"AI tạm thời không phản hồi: {detail}") from None
        raise ExcelAIUnavailable("AI tạm thời không phản hồi. Vui lòng thử lại sau.") from None
    except Exception as exc:
        raise ExcelAIUnavailable("Không kết nối được AI. Vui lòng thử lại sau.") from exc
    if len(raw) > 100000:
        raise ExcelAIUnavailable("AI trả lời quá dài. Hãy chọn vùng nhỏ hơn.")
    try:
        reply = json.loads(raw)
    except ValueError as exc:
        raise ExcelAIUnavailable("AI trả về dữ liệu không hợp lệ.") from exc
    text = reply.get("text")
    if not isinstance(text, str) or not text.strip() or len(text) > 12000:
        raise ExcelAIUnavailable("AI trả về câu trả lời không hợp lệ.")
    return text


def explain(
    workbook: ExcelWorkbook,
    user_id: int,
    sheet_name: str,
    scope: str,
    selection: str,
    question: str,
    history: list[dict],
) -> dict[str, Any]:
    context, cells_examined, references_examined, truncated = build_context(workbook, sheet_name, scope, selection)
    address = context["selection"]["address"]
    prompt = (
        f"Giải thích thật chi tiết vùng {address} trong file {workbook.original_filename}. "
        "Nêu rõ: công thức và ý nghĩa từng thành phần; nguồn số liệu/ô tham chiếu; cách kiểm tra chéo; "
        "các lỗi hoặc giả định tiềm ẩn. Phân biệt công thức với giá trị cache; chỉ dùng dữ liệu trong context."
    )
    if question:
        prompt += " Câu hỏi thêm của người dùng: " + question
    explanation = _gateway_reply(context, prompt, history, user_id, workbook.code)
    return {
        "explanation": explanation,
        "cells_examined": cells_examined,
        "references_examined": references_examined,
        "truncated": truncated,
    }
